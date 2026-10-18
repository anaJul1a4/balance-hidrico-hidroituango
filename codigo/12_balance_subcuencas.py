"""
Punto 9 - Validación del balance hídrico de largo plazo en las áreas de drenaje de las estaciones de caudal.

  P = ETR + R + ΔS,  ΔS ≈ 0  ⇒  R = P − ETR ;   Q_est = R · A / 31 557 600 s   (R en m, A en m²)
  P: kriging IDEAM 1981-2010 (punto 3) · ETP: Hargreaves (punto 6) · ETR: Budyko celda a celda (punto 7)
  Q_obs: caudal medio anual multianual 1981-2010 (punto 8; años completos) y serie rellenada (complemento).

Unidades de análisis: área de drenaje de cada estación (anidadas a lo largo del Cauca), tramos intermedios entre
estaciones consecutivas y la cuenca hasta Hidroituango (Q_obs transferido desde Pto. Valdivia por áreas y
Pte. Pescadero rellenada).
Sensibilidad: P por IDW, P con todas las estaciones, ETP escalada a la ETP del IDEAM, Budyko agregado.

Salida: trabajo1/resultados/09_balance/Punto9_balance_hidrico.xlsx (+ figuras)
"""
from pathlib import Path

import geopandas as gpd
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import rasterio
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from rasterio.features import geometry_mask

ROOT = Path(__file__).resolve().parents[2]
RES = ROOT / "trabajo1" / "resultados"
DAT = ROOT / "trabajo1" / "datos" / "caudales_ideam"
OUT = RES / "09_balance"
OUT.mkdir(parents=True, exist_ok=True)
SEG = 365.25 * 86400
P0, P1 = "1981", "2010"
CIERRE = "HIDRO"


def leer(p):
    with rasterio.open(p) as s:
        a = s.read(1).astype("float64")
        a[a == s.nodata] = np.nan
        return a, s.transform


P, aff = leer(RES / "04_precipitacion_interpolada" / "P_anual_kriging_90m_3116.tif")
ETP, _ = leer(RES / "08_etp_etr" / "ETP_anual_hargreaves_90m_3116.tif")
ESC_P = {"P kriging (base)": P,
         "P IDW (p=1.5)": leer(RES / "04_precipitacion_interpolada" / "P_anual_idw_90m_3116.tif")[0],
         "P todas las estaciones (IDW p=2.5)": leer(RES / "04_precipitacion_interpolada" / "todas_estaciones" / "P_anual_idw_90m_3116.tif")[0]}
k_ide = pd.read_csv(RES / "08_etp_etr" / "contraste_ETP_IDEAM.csv", index_col=0).iloc[:, 0]
K_ETP = k_ide["ETP IDEAM media (mm)"] / k_ide["ETP Hargreaves en esos puntos (mm)"]


def budyko(p, etp):
    f = etp / p
    return p * np.sqrt(f * np.tanh(1 / f) * (1 - np.exp(-f)))


ESC = {k: (p, ETP) for k, p in ESC_P.items()}
ESC[f"ETP escalada a IDEAM (×{K_ETP:.3f})"] = (P, ETP * K_ETP)

# ----------------------------------------------------------------------------- unidades de análisis
sub = gpd.read_file(RES / "06_caudales" / "subcuencas_estaciones.gpkg", layer="areas_drenaje")
cuenca = gpd.read_file(RES / "01_delimitacion" / "cuenca_cauca_hidroituango.gpkg", layer="cuenca")
sub = sub[sub.codigo != "26237040"].copy()                      # Pto. Valdivia drena fuera de los rásters
sub = pd.concat([sub, gpd.GeoDataFrame({"codigo": [CIERRE], "estacion": ["HIDROITUANGO (cierre)"],
                                        "area_km2_DEM": [float(cuenca.area_km2[0])]}, geometry=[cuenca.geometry[0]],
                                       crs=sub.crs)]).sort_values("area_km2_DEM").reset_index(drop=True)
masks = {r.codigo: ~geometry_mask([r.geometry], out_shape=P.shape, transform=aff) for r in sub.itertuples()}
ca = abs(aff.a * aff.e) / 1e6

# caudales observados
Qm = pd.read_csv(DAT / "Q_mensual_completa.csv", index_col=0, parse_dates=True).loc[P0:P1]
Qr = pd.read_csv(DAT / "Q_mensual_1981_2010_rellenado.csv", index_col=0, parse_dates=True)
anual = lambda s: s.groupby(s.index.year).apply(lambda x: x.mean() if x.count() == 12 else np.nan)
Qobs = {c: anual(Qm[c]).mean() for c in Qm.columns if c in set(sub.codigo) | {"26237040"}}
Qrel = {c: anual(Qr[c]).mean() for c in Qr.columns}
A_V = 37267.0
A_C = float(cuenca.area_km2[0])
Qobs[CIERRE] = Qobs["26237040"] * A_C / A_V
Qrel[CIERRE] = Qrel["26237050"] * A_C / 35450.77           # Pte. Pescadero rellenada (98.5 % del área)

# ----------------------------------------------------------------------------- estadísticas zonales
filas, sens = [], []
for r in sub.itertuples():
    m = masks[r.codigo]
    A = m.sum() * ca
    p, e = np.nanmean(P[m]), np.nanmean(ETP[m])
    etr = np.nanmean(budyko(P[m], ETP[m]))
    filas.append({"codigo": r.codigo, "estacion": r.estacion.title() if r.codigo != CIERRE else r.estacion,
                  "A (km²)": A, "P (mm/año)": p, "ETP (mm/año)": e, "ETR Budyko celda a celda (mm/año)": etr,
                  "ETR Budyko agregado (mm/año)": budyko(p, e),
                  "Q obs (m³/s)": Qobs.get(r.codigo, np.nan), "Q obs rellenado (m³/s)": Qrel.get(r.codigo, np.nan)})
    for k, (pp, ee) in ESC.items():
        sens.append({"codigo": r.codigo, "escenario": k, "P": np.nanmean(pp[m]), "ETP": np.nanmean(ee[m]),
                     "ETR": np.nanmean(budyko(pp[m], ee[m])), "A": A})
bal = pd.DataFrame(filas)
sens = pd.DataFrame(sens)
sens["R"] = sens.P - sens.ETR
sens["Q_est"] = sens.R / 1000 * sens.A * 1e6 / SEG
sens = sens.merge(bal[["codigo", "Q obs (m³/s)"]], on="codigo")
sens["error (%)"] = 100 * (sens.Q_est / sens["Q obs (m³/s)"] - 1)
b = bal.copy()
b["R est (mm)"] = b["P (mm/año)"] - b["ETR Budyko celda a celda (mm/año)"]
b["Q est"] = b["R est (mm)"] / 1000 * b["A (km²)"] * 1e6 / SEG
b["error %"] = 100 * (b["Q est"] / b["Q obs (m³/s)"] - 1)
b["R obs (mm)"] = b["Q obs (m³/s)"] * SEG / (b["A (km²)"] * 1e6) * 1000
pd.set_option("display.width", 250)
print(b.round(1).to_string())
print(sens.pivot(index="codigo", columns="escenario", values="error (%)").round(1).to_string())

# ----------------------------------------------------------------------------- tramos intermedios
orden = list(b.codigo)
tr = []
for aguas_arriba, aguas_abajo in zip([None] + orden[:-1], orden):
    d, u = b.set_index("codigo").loc[aguas_abajo], (b.set_index("codigo").loc[aguas_arriba] if aguas_arriba else None)
    tr.append({"tramo": (f"{u.estacion} → {d.estacion}" if aguas_arriba else f"Cabecera → {d.estacion}"),
               "abajo": aguas_abajo, "arriba": aguas_arriba})
tr = pd.DataFrame(tr)

# ----------------------------------------------------------------------------- figuras
fig, axs = plt.subplots(1, 2, figsize=(13, 5.5))
ax = axs[0]
ax.scatter(b["Q obs (m³/s)"], b["Q est"], s=50, c="#2171b5", edgecolors="black", zorder=3, label="Q obs (años completos)")
ax.scatter(b["Q obs rellenado (m³/s)"], b["Q est"], s=25, marker="s", facecolors="none", edgecolors="#d73027", zorder=3,
           label="Q obs rellenado")
for _, r in b.iterrows():
    ax.annotate(r.estacion.replace(" (cierre)", ""), (r["Q obs (m³/s)"], r["Q est"]), xytext=(5, -3),
                textcoords="offset points", fontsize=7)
lim = [0, 1250]
ax.plot(lim, lim, "k-", lw=0.8); ax.fill_between(lim, np.array(lim) * 0.9, np.array(lim) * 1.1, color="0.9", label="±10 %")
ax.set_xlim(lim); ax.set_ylim(lim); ax.set_xlabel("Q observado (m³/s)"); ax.set_ylabel("Q estimado = (P − ETR)·A (m³/s)")
ax.set_title("Balance de largo plazo: Q estimado vs. observado (1981–2010)", fontsize=10); ax.legend(fontsize=8); ax.grid(alpha=0.3)
ax = axs[1]
f = np.linspace(0.05, 2, 200)
ax.plot(f, np.sqrt(f * np.tanh(1 / f) * (1 - np.exp(-f))), "k", label="Budyko (1974)")
ax.plot([0, 1, 2], [0, 1, 1], "k--", lw=0.7)
phi = b["ETP (mm/año)"] / b["P (mm/año)"]
ax.scatter(phi, 1 - b["R obs (mm)"] / b["P (mm/año)"], s=50, c="#d73027", edgecolors="black", zorder=3,
           label="Observado: ETR = P − Q_obs")
ax.scatter(phi, b["ETR Budyko celda a celda (mm/año)"] / b["P (mm/año)"], s=30, marker="D", c="#2171b5", zorder=3,
           label="Budyko celda a celda")
for (x, y, n) in zip(phi, 1 - b["R obs (mm)"] / b["P (mm/año)"], b.estacion):
    ax.annotate(n.replace(" (cierre)", ""), (x, y), xytext=(5, 2), textcoords="offset points", fontsize=7)
ax.set_xlim(0.4, 1.2); ax.set_ylim(0.3, 0.85); ax.set_xlabel("φ = ETP / P (media de la subcuenca)"); ax.set_ylabel("ETR / P")
ax.set_title("Subcuencas en el espacio de Budyko", fontsize=10); ax.legend(fontsize=8, loc="lower right"); ax.grid(alpha=0.3)
fig.tight_layout(); fig.savefig(OUT / "balance_Q_est_vs_obs.png", dpi=200); plt.close(fig)

# ----------------------------------------------------------------------------- Excel
wb = Workbook()
H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F4E79"); BOLD = Font(bold=True)
thin = Side(style="thin", color="BBBBBB"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
INP = PatternFill("solid", fgColor="FFF2CC"); OK = PatternFill("solid", fgColor="C6EFCE"); MAL = PatternFill("solid", fgColor="FFC7CE")


def encabezado(ws, fila, cols, alto=48):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(fila, j, c); cell.font = H; cell.fill = HF; cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.row_dimensions[fila].height = alto


ws = wb.active; ws.title = "Metodologia"
for i, (s, f) in enumerate([
        ("PUNTO 9 — VALIDACIÓN DEL BALANCE HÍDRICO DE LARGO PLAZO", BOLD),
        ("P = ETR + R + ΔS ;  para periodos largos ΔS ≈ 0  ⇒  R = P − ETR ;  Q_estimado = R·A  (R en m/año, A en m², ÷ 31 557 600 s)", None),
        ("P: precipitación media anual 1981–2010, kriging ordinario de estaciones IDEAM (punto 3).", None),
        ("ETP: Hargreaves-Samani mensual con Ra por latitud de cada celda (punto 6).  ETR: Budyko (1974) celda a celda (punto 7).", None),
        ("P, ETP y ETR son promedios de las celdas de 90 m dentro del área de drenaje de cada estación (datos de entrada, amarillo).", None),
        ("Q_obs: caudal medio anual multianual 1981–2010 de años completos (punto 8). Q_obs rellenado: serie rellenada por regresión entre estaciones.", None),
        ("Cierre (Hidroituango): Q_obs transferido desde Pto. Valdivia por relación de áreas (37 267 → 35 992 km²); "
         "Q_obs rellenado desde Pte. Pescadero (35 451 → 35 992 km²).", None),
        ("Pto. Valdivia no se evalúa directamente porque su área de drenaje sale de la cuenca delimitada (no hay P/ETP fuera de ella).", None),
        ("Tramos: área incremental entre estaciones consecutivas del Cauca (P, ETR y Q por diferencia de volúmenes).", None),
        ("Error = (Q_est − Q_obs) / Q_obs. Verde: |error| ≤ 10 %; rojo: |error| > 20 %.", None)], 1):
    c = ws.cell(i, 1, s)
    if f:
        c.font = f
ws.column_dimensions["A"].width = 150

# Balance por área de drenaje (fórmulas)
ws = wb.create_sheet("Balance_subcuencas")
cols = ["Código", "Estación", "A (km²)", "P (mm/año)", "ETP (mm/año)", "φ = ETP/P", "ETR Budyko (mm/año)", "ETR/P",
        "R = P − ETR (mm/año)", "Q estimado (m³/s)", "Q observado (m³/s)", "R observado (mm/año)", "ETR implícita = P − R obs (mm/año)",
        "Error Q (%)", "Q obs rellenado (m³/s)", "Error vs. rellenado (%)", "ETR Budyko agregado (mm/año)", "Error Budyko agregado (%)"]
encabezado(ws, 1, cols)
n = len(b)
for i, (_, r) in enumerate(b.iterrows(), 2):
    vals = [r.codigo, r.estacion, round(r["A (km²)"], 1), round(r["P (mm/año)"], 1), round(r["ETP (mm/año)"], 1),
            f"=E{i}/D{i}", round(r["ETR Budyko celda a celda (mm/año)"], 1), f"=G{i}/D{i}",
            f"=D{i}-G{i}", f"=I{i}/1000*C{i}*1000000/{SEG:.0f}", round(r["Q obs (m³/s)"], 2),
            f"=K{i}*{SEG:.0f}/(C{i}*1000000)*1000", f"=D{i}-L{i}", f"=(J{i}/K{i}-1)*100",
            round(r["Q obs rellenado (m³/s)"], 2), f"=(J{i}/O{i}-1)*100", round(r["ETR Budyko agregado (mm/año)"], 1),
            f"=((D{i}-Q{i})/1000*C{i}*1000000/{SEG:.0f}/K{i}-1)*100"]
    for j, x in enumerate(vals, 1):
        c = ws.cell(i, j, x); c.border = BOX
        c.number_format = "0.00" if j in (6, 8) else ("0.0" if j > 2 else "General")
        if j in (3, 4, 5, 7, 11, 15, 17):
            c.fill = INP
    e = abs(r["error %"])
    ws.cell(i, 14).fill = OK if e <= 10 else (MAL if e > 20 else PatternFill())
fi = n + 3
ws.cell(fi, 1, "Medidas de desempeño (estaciones, sin el cierre)").font = BOLD
rng = f"N2:N{n}"
for k, (lab, f) in enumerate([("Sesgo medio (%)", f"=AVERAGE({rng})"), ("Error absoluto medio (%)", f"=SUMPRODUCT(ABS({rng}))/COUNT({rng})"),
                              ("RMSE relativo (%)", f"=SQRT(SUMSQ({rng})/COUNT({rng}))"),
                              ("Correlación Q est vs Q obs", f"=CORREL(J2:J{n},K2:K{n})")], 1):
    ws.cell(fi + k, 1, lab).border = BOX
    c = ws.cell(fi + k, 2, f); c.border = BOX; c.number_format = "0.00"
for j, w in enumerate([9, 22, 9, 9, 9, 8, 10, 7, 10, 10, 10, 10, 12, 9, 10, 10, 11, 10], 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "C2"

# Tramos intermedios (fórmulas sobre Balance_subcuencas)
ws = wb.create_sheet("Tramos")
encabezado(ws, 1, ["Tramo", "A incremental (km²)", "P (mm/año)", "ETR (mm/año)", "R estimado (mm/año)", "Q estimado incremental (m³/s)",
                   "Q observado incremental (m³/s)", "R observado (mm/año)", "Error Q incremental (%)",
                   "Q obs incremental, serie rellenada (mismos 30 años) (m³/s)", "Error vs. rellenado (%)", "Observación"])
pos = {c: i + 2 for i, c in enumerate(b.codigo)}
notas = {"26027010": "Regulado por el embalse de Salvajina (desde 1985)",
         "26107070": "Valle del Cauca: riego y extracciones; cauce del DEM a 3–5 km del real",
         "26067010": "Incluye ríos del sur del Valle; Juanchito ajustado 3.4 km"}
for i, t in enumerate(tr.itertuples(), 2):
    d = pos[t.abajo]
    if pd.isna(t.arriba):
        f = [f"=Balance_subcuencas!C{d}", f"=Balance_subcuencas!D{d}", f"=Balance_subcuencas!G{d}", f"=C{i}-D{i}",
             f"=Balance_subcuencas!J{d}", f"=Balance_subcuencas!K{d}"]
    else:
        u = pos[t.arriba]
        f = [f"=Balance_subcuencas!C{d}-Balance_subcuencas!C{u}",
             f"=(Balance_subcuencas!D{d}*Balance_subcuencas!C{d}-Balance_subcuencas!D{u}*Balance_subcuencas!C{u})/B{i}",
             f"=(Balance_subcuencas!G{d}*Balance_subcuencas!C{d}-Balance_subcuencas!G{u}*Balance_subcuencas!C{u})/B{i}",
             f"=C{i}-D{i}", f"=Balance_subcuencas!J{d}-Balance_subcuencas!J{u}", f"=Balance_subcuencas!K{d}-Balance_subcuencas!K{u}"]
    qrel = f"=Balance_subcuencas!O{d}" if pd.isna(t.arriba) else f"=Balance_subcuencas!O{d}-Balance_subcuencas!O{pos[t.arriba]}"
    vals = [t.tramo] + f + [f"=G{i}*{SEG:.0f}/(B{i}*1000000)*1000", f"=(F{i}/G{i}-1)*100", qrel, f"=(F{i}/J{i}-1)*100",
                            notas.get(t.abajo, "")]
    for j, x in enumerate(vals, 1):
        c = ws.cell(i, j, x); c.border = BOX; c.number_format = "0.0"
ft = len(tr) + 3
ws.cell(ft, 1, "Nota: el caudal incremental es la diferencia de dos caudales grandes, cada uno con error de medición de "
        "3–5 % (25–45 m³/s), comparable al propio caudal del tramo. Con 'años completos' cada estación promedia años "
        "distintos; la serie rellenada usa los mismos 30 años en todas y da incrementos más consistentes.")
for j, w in enumerate([40, 11, 10, 10, 10, 12, 12, 11, 11, 14, 11, 55], 1):
    ws.column_dimensions[get_column_letter(j)].width = w

# Sensibilidad
ws = wb.create_sheet("Sensibilidad")
piv = sens.pivot(index="codigo", columns="escenario", values="Q_est").loc[b.codigo]
pe = sens.pivot(index="codigo", columns="escenario", values="error (%)").loc[b.codigo]
escs = list(ESC)
encabezado(ws, 1, ["Código", "Estación", "Q observado (m³/s)"] + [f"Q est: {e}" for e in escs] + [f"Error: {e} (%)" for e in escs])
for i, c in enumerate(b.codigo, 2):
    vals = [c, b.set_index("codigo").loc[c, "estacion"], round(b.set_index("codigo").loc[c, "Q obs (m³/s)"], 1)] + \
           [round(piv.loc[c, e], 1) for e in escs] + [round(pe.loc[c, e], 1) for e in escs]
    for j, x in enumerate(vals, 1):
        ws.cell(i, j, x).border = BOX
r0 = len(b) + 3
ws.cell(r0, 1, "Cierre (Hidroituango): rango de Q observado según la fuente").font = BOLD
for k, (lab, v) in enumerate([("Pto. Valdivia, años completos, transferido por áreas", Qobs[CIERRE]),
                              ("Pto. Valdivia, serie rellenada, transferido", Qrel["26237040"] * A_C / A_V),
                              ("Pte. Pescadero, serie rellenada, transferido", Qrel[CIERRE])], 1):
    ws.cell(r0 + k, 1, lab).border = BOX; ws.cell(r0 + k, 2, round(v, 1)).border = BOX
    ws.cell(r0 + k, 3, f"=({round(piv.loc[CIERRE, escs[0]], 2)}/B{r0 + k}-1)*100").number_format = "0.0"
ws.cell(r0, 3, "Error Q est base (%)").font = BOLD
for j in range(1, 3 + 2 * len(escs) + 1):
    ws.column_dimensions[get_column_letter(j)].width = 14 if j > 2 else (12 if j == 1 else 22)

# Incertidumbre
ws = wb.create_sheet("Incertidumbre")
txt = [
    ("FUENTES DE INCERTIDUMBRE", BOLD),
    ("Precipitación", BOLD),
    ("· Densidad de estaciones muy baja sobre 3000 m (páramos, nevados): el kriging tiende a subestimar la lluvia orográfica de alta montaña.", None),
    ("· Elección del vecindario: incluir estaciones de otras vertientes cambia P en los bordes (ver Sensibilidad).", None),
    ("· Error de interpolación (validación cruzada): RMSE ≈ 363 mm por estación; desviación estándar del kriging ≈ 300 mm.", None),
    ("Evapotranspiración", BOLD),
    (f"· Hargreaves sobreestima la ETP del IDEAM en ≈ 22 % (factor {K_ETP:.3f}); Budyko traslada parte de ese sesgo a la ETR.", None),
    ("· Temperaturas derivadas de una única regresión T–Z para toda la cuenca (no capta inversiones ni efecto de exposición).", None),
    ("· Budyko es una relación empírica de largo plazo sin parámetro de calibración (no considera vegetación, suelos ni estacionalidad).", None),
    ("· Budyko aplicado de forma agregada (φ medio) subestima R frente a la aplicación celda a celda, por la no linealidad de la curva.", None),
    ("Caudal", BOLD),
    ("· Curvas de calibración nivel–caudal del IDEAM (errores típicos 5–15 % en crecientes y estiaje).", None),
    ("· Faltantes: 10–22 % de los días en Pto. Valdivia y estaciones del Valle; las medias varían ≤ 5 % según se rellenen o no.", None),
    ("· Regulación (Salvajina desde 1985) y extracciones para riego y acueductos en el Valle del Cauca: el caudal observado no es natural.", None),
    ("· Cierre: no hay estación validada en la presa; se transfiere desde Pto. Valdivia (área 3.5 % mayor) o Pte. Pescadero (rellenada).", None),
    ("Interpolación espacial y delimitación", BOLD),
    ("· Áreas de drenaje con DEM de 90 m; en el Valle el cauce del DEM se separa 3–5 km del real (±5 % de área en Juanchito y La Victoria).", None),
    ("Periodos diferentes", BOLD),
    ("· P y T son promedios 1981–2010 de estaciones con registros de distinta longitud; los caudales usan sólo años completos dentro de 1981–2010.", None),
    ("· El supuesto ΔS ≈ 0 es razonable a 30 años, pero el embalse de Salvajina y los acuíferos del Valle pueden almacenar o perder agua.", None),
]
for i, (s, f) in enumerate(txt, 1):
    c = ws.cell(i, 1, s)
    if f:
        c.font = f
ws.column_dimensions["A"].width = 150

ws = wb.create_sheet("Figuras")
img = XLImage(str(OUT / "balance_Q_est_vs_obs.png")); img.width, img.height = 1040, 440
ws.add_image(img, "A1")
img = XLImage(str(RES / "08_etp_etr" / "mapas_ETP_phi_ETR_R_panel.png")); img.width, img.height = 1100, 606
ws.add_image(img, "A25")
wb.save(OUT / "Punto9_balance_hidrico.xlsx")
b.to_csv(OUT / "balance_subcuencas.csv", index=False, encoding="utf-8-sig")
print("Excel guardado")
