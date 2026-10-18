"""
Documento de cambios para el informe del equipo («Trabajo 1 Hidrometeorología David y Juli»): lista, en el orden del informe, lo
que hay que corregir, reemplazar o agregar, con el texto nuevo resaltado en amarillo y las figuras listas para pegar. La numeración
de secciones, figuras y tablas es la de esa versión (12 secciones, 16 tablas, 20 figuras).

Salida: ~/Downloads/Cambios para el informe - David y Juli.docx
"""
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RES = Path(__file__).resolve().parents[1] / "resultados"
FIG = RES / "informe_figs"
SALIDA = Path.home() / "Downloads" / "Cambios para el informe - David y Juli.docx"

d = docx.Document()
s = d.sections[0]
s.left_margin = s.right_margin = Cm(2.5); s.top_margin = s.bottom_margin = Cm(2.2)
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
for n in ("Heading 1", "Heading 2"):
    h = d.styles[n]; h.font.name = "Times New Roman"; h.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def P(texto, amarillo=False, cursiva=False, gris=False, negrita=False, centro=False, tam=None):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(texto); r.italic = cursiva; r.bold = negrita
    if amarillo:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    if gris:
        r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    if tam:
        r.font.size = Pt(tam)
    return p


def cambio(n, titulo, donde, actual=None, nuevo=None, nota=None):
    d.add_heading(f"{n}. {titulo}", 2)
    p = d.add_paragraph(); p.add_run("Dónde: ").bold = True; p.add_run(donde)
    if actual:
        p = d.add_paragraph(); p.add_run("Texto actual: ").bold = True
        r = p.add_run(actual); r.italic = True; r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    if nuevo:
        d.add_paragraph().add_run("Texto nuevo:").bold = True
        for par in nuevo if isinstance(nuevo, list) else [nuevo]:
            P(par, amarillo=True)
    if nota:
        P(nota, cursiva=True, gris=True, tam=10)


def figura(nombre, ancho_cm, leyenda):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIG / nombre), width=Cm(ancho_cm))
    q = d.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    n, resto = leyenda.split(". ", 1)
    a = q.add_run(n + ". "); a.bold = True; a.font.size = Pt(9); a.font.highlight_color = WD_COLOR_INDEX.YELLOW
    b = q.add_run(resto); b.font.size = Pt(9); b.font.highlight_color = WD_COLOR_INDEX.YELLOW


def tabla(enc, filas, anchos):
    t = d.add_table(rows=1, cols=len(enc)); t.style = "Table Grid"
    for j, e in enumerate(enc):
        c = t.rows[0].cells[j]; c.text = ""; r = c.paragraphs[0].add_run(e); r.bold = True; r.font.size = Pt(8.5)
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    for f in filas:
        cs = t.add_row().cells
        for j, v in enumerate(f):
            cs[j].text = ""; r = cs[j].paragraphs[0].add_run(str(v)); r.font.size = Pt(8.5)
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            cs[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
    for row in t.rows:
        for j, w in enumerate(anchos):
            row.cells[j].width = Cm(w)


# ============================================================================= portada del documento
P("Cambios para el informe «Trabajo 1 Hidrometeorología David y Juli»", negrita=True, centro=True, tam=14)
P("Lo que hay que copiar al documento original está resaltado en amarillo. La numeración de secciones, figuras y tablas es la de "
  "esa versión. Los cambios van en el orden del informe: primero correcciones pequeñas, luego las figuras con la nueva escala de "
  "color, la subsección nueva sobre la corrección del sesgo de ERA5 (que entra al final de la sección 12, para no renumerar nada) "
  "y por último la conclusión y la referencia.", cursiva=True, gris=True, tam=10)

d.add_heading("A. Correcciones de texto", 1)
cambio(1, "Área de la delimitación con GLO-30", "Sección 1, párrafo después de la Figura 1.",
       "Para los cálculos de las secciones siguientes se usó una delimitación hecha con el DEM Copernicus GLO-30, su área coincide con "
       "la de QGIS (35992 km²) con una diferencia de 0.4%.",
       "Para los cálculos de las secciones siguientes se usó una delimitación hecha con el DEM Copernicus GLO-30 remuestreado a 90 m, "
       "que da 35 992 km². Su área coincide con la de HydroSHEDS (35 841 km²; Lehner et al., 2008) con una diferencia de 0.4 %. La "
       "diferencia con la Tabla 1 (35 283 km², 2 %) se debe a la resolución del DEM usado en QGIS.",
       "35 992 km² es el área de GLO-30, no la de QGIS; la de QGIS es la de la Tabla 1. El 0.4 % es la diferencia con HydroSHEDS.")
cambio(2, "Explicación del sesgo de ERA5", "Sección 9, final del párrafo que empieza «Cada producto se convirtió a mm/mes…».",
       "Parte de la sobreestimación de la precipitación de ERA 5 y ERA 5 LAND podría ser explicada al considerar que esta se acumula a "
       "partir de los pronósticos de corto plazo del modelo, es decir, no hay una corrección de la lluvia combinando el satélite con "
       "estaciones. Otra posible explicación podría ser la resolución espacial y la topografía compleja que el modelo suaviza y no "
       "representa bien.",
       "Las unidades se verificaron: en la Amazonia y los Llanos ERA5 coincide con CHIRPS (razón de 0.89 a 0.98), y la sobrestimación "
       "crece en los Andes hasta ≈ 3 en esta cuenca. La precipitación de ERA5 no es una observación: se acumula a partir de los "
       "pronósticos de corto plazo del modelo, y ERA5 no asimila pluviómetros (Hersbach et al., 2020); ERA5-Land la toma de ERA5 por "
       "interpolación (Muñoz-Sabater et al., 2021). La resolución, en cambio, no explica el exceso: el campo del IDEAM degradado a la "
       "celda de ERA5 (0.25°) conserva su media y su balance sigue cerrando (−5.6 %), y remuestrear ERA5 a 90 m de forma bilineal no "
       "cambia su error (+391 %). El sesgo surge de la física del modelo sobre una topografía mal representada (dos cordilleras y un "
       "valle interandino) y persiste porque no hay estaciones que lo corrijan.",
       "«combinando el satélite con estaciones» describe a CHIRPS, no a ERA5. El nuevo texto agrega las dos pruebas que hicimos "
       "(unidades y resolución).")
cambio(3, "Referencia a una tabla que no existe", "Sección 9, párrafo antes de la Figura 13.",
       "La Tabla 15 resume sus estadísticos sobre la cuenca, y la Figura 14 muestra la razón entre cada producto y el campo del IDEAM.",
       "Sobre la cuenca, la precipitación media es de 1985 mm/año con el IDEAM, 2026 mm/año con CHIRPS (+2 %), 5961 mm/año con "
       "ERA5-Land y 5976 mm/año con ERA5 (+200 %); la correlación espacial con el IDEAM es de 0.89 para CHIRPS y de −0.06 para los dos "
       "reanálisis. La Figura 14 muestra la razón entre cada producto y el campo del IDEAM.",
       "La Tabla 15 del informe es la de precipitación mensual; la tabla de estadísticos por producto no quedó en esta versión.")
cambio(4, "Título de la sección 10", "Título de la sección 10.", "Blande hídrico con cada producto", "Balance hídrico con cada producto")
cambio(5, "Referencia a la sección del balance", "Sección 10, primera frase.", "…se repitió el balance del punto 9…",
       "…se repitió el balance de la sección 8…")
cambio(6, "Errata", "Sección 11, primer párrafo.", "…según quedará por encima o por debajo…", "…según quedara por encima o por debajo…")
cambio(7, "Figura citada", "Sección 11, segundo párrafo.", "En cada celda se identificaron además tres cosas (Figura 16):",
       "En cada celda se identificaron además tres cosas (Figura 18):")
cambio(8, "Leyenda de la Figura 18 incompleta", "Leyenda de la Figura 18.",
       "Figura 18. Mes de máxima lluvia, número de picos del ciclo anual e índice de estacionalidad",
       "Figura 18. Mes de máxima lluvia, número de picos del ciclo anual e índice de estacionalidad de Walsh y Lawler en cada celda.")

d.add_heading("B. Figuras de precipitación con la nueva escala de color", 1)
P("Todos los mapas usan ahora un mismo código de color: para el agua disponible (precipitación y escorrentía), rojo = poca agua, "
  "amarillo cerca de la media de la cuenca y azul = mucha agua; para la evaporación (ETP y ETR), blanco → amarillo → rojo, con las "
  "mismas clases en los dos mapas. Ya no hay verdes. Cada valor tiene el mismo color en todos los mapas de su variable, incluida la "
  "figura nueva del mapeo de cuantiles. Reemplazar la imagen de las Figuras 2, 4, 8, 13, 15 y 16 por estas.",
  cursiva=True, gris=True, tam=10)
for nombre, ancho, ley in [("mapa_estaciones_P.png", 8.5, "Figura 2. (reemplazar imagen; la leyenda no cambia)"),
                           ("mapa_P_anual.png", 8.5, "Figura 4. (reemplazar imagen; la leyenda no cambia)"),
                           ("panel_ETP_phi_ETR_R.png", 16, "Figura 8. Evapotranspiración potencial de Hargreaves, índice de aridez φ, "
                            "evapotranspiración real de Budyko y escorrentía R = P − ETR. ETP y ETR comparten paleta y clases; la "
                            "escorrentía usa el código de la precipitación (rojo = poca agua, azul = mucha)."),
                           ("panel_P_productos.png", 16, "Figura 13. (reemplazar imagen; la leyenda no cambia)"),
                           ("panel_R_productos.png", 16, "Figura 15. (reemplazar imagen; la leyenda no cambia)"),
                           ("panel_P_mensual.png", 16, "Figura 16. (reemplazar imagen; la leyenda no cambia)")]:
    figura(nombre, ancho, ley)

d.add_heading("C. Subsección nueva: corrección del sesgo de ERA5 y ERA5-Land", 1)
P("Dónde: al final de la sección 12, después del párrafo «En conjunto, CHIRPS es el mejor producto…». Va como subsección con "
  "título en negrita, igual que «Diagrama de Taylor».", cursiva=True, gris=True, tam=10)
P("Corrección del sesgo de ERA5 y ERA5-Land por mapeo de cuantiles (análisis adicional)", amarillo=True, negrita=True)
for t in [
    "Para saber si los reanálisis pueden usarse una vez corregidos, se aplicó un mapeo de cuantiles empírico (Gudmundsson et al., "
    "2012), celda a celda y por mes calendario: cada valor mensual de ERA5 se reemplaza por el valor de la referencia que ocupa el "
    "mismo cuantil en la distribución de ese mes. Del IDEAM solo se tienen las normales 1981–2010, no series mensuales por estación, "
    "así que se usó como referencia CHIRPS v2.0, que está corregido con estaciones, promediado en la celda de cada reanálisis. El "
    "método corrige toda la distribución (media, varianza y extremos), pero conserva el orden de los valores de ERA5 en cada mes, así "
    "que no corrige el momento en que llueve. Para evaluarlo sin circularidad se calibró con 1981–1995 y se validó con 1996–2010, y "
    "el balance se comparó con los caudales observados, que el método no usa.",
    "La Figura 21 muestra la precipitación y la escorrentía de los dos reanálisis antes y después de la corrección, con las mismas "
    "escalas de color de los demás mapas, y la Tabla 17 resume el balance y la variabilidad."]:
    P(t, amarillo=True)
figura("panel_QM_P_R.png", 16, "Figura 21. Precipitación media anual (arriba) y escorrentía de largo plazo R = P − ETR (abajo) de ERA5 y "
       "ERA5-Land sin corregir y corregidas por mapeo de cuantiles (1981–2010), con la ETP de Hargreaves y Budyko celda a celda.")
p = d.add_paragraph(); a = p.add_run("Tabla 17. "); a.bold = True; a.italic = True; a.font.size = Pt(9)
a.font.highlight_color = WD_COLOR_INDEX.YELLOW
b = p.add_run("Balance y variabilidad de la precipitación de ERA5 y ERA5-Land antes y después del mapeo de cuantiles (1981–2010).")
b.italic = True; b.font.size = Pt(9); b.font.highlight_color = WD_COLOR_INDEX.YELLOW
tabla(["Fuente", "P (mm/año)", "Error en el cierre (%)", "Error abs. medio, 9 estaciones (%)", "r espacial con el IDEAM",
       "CV interanual (%)", "r anual P–Q (Bolombolo / La Pintada / Pto. Valdivia)"],
      [["IDEAM (kriging)", "1985", "−5.3", "4.5", "—", "—", "—"],
       ["CHIRPS v2.0", "2026", "−2.2", "5.3", "0.89", "10.8", "0.92 / 0.94 / 0.91"],
       ["ERA5 sin corregir", "5976", "+392", "447", "−0.06", "7.1", "0.56 / 0.58 / 0.78"],
       ["ERA5 corregido", "2057", "+0.5", "6.7", "0.81", "8.7", "0.62 / 0.64 / 0.82"],
       ["ERA5-Land sin corregir", "5961", "+390", "452", "−0.06", "7.6", "0.58 / 0.59 / 0.79"],
       ["ERA5-Land corregido", "2020", "−2.7", "5.2", "0.88", "9.6", "0.62 / 0.64 / 0.82"]],
      [3.0, 1.6, 1.8, 2.2, 1.8, 1.6, 3.6])
for t in [
    "En la validación (1996–2010, años no usados en la calibración), el sesgo frente a CHIRPS pasa de +184 % a −2.7 % en ERA5 y de "
    "+193 % a −1.5 % en ERA5-Land, y la correlación mensual sube de 0.71–0.76 a 0.91. Con la corrección completa, el balance cierra "
    "con +0.5 % (ERA5) y −2.7 % (ERA5-Land), el error en las estaciones queda en 5–7 %, como con el IDEAM y CHIRPS, el ciclo anual "
    "coincide con el del IDEAM (máximo en octubre) y el patrón espacial se recupera (r = 0.81 y 0.88). La escorrentía media pasa de "
    "unos 4670 mm/año a 924–955 mm/año, cerca de los 950 mm/año observados (Figura 21).",
    "Sin embargo, casi todo lo que se corrige viene de la referencia: la media, el ciclo y el patrón espacial corregidos son, en "
    "esencia, los de CHIRPS vistos con la celda del reanálisis. La única información propia de ERA5 que sobrevive es la secuencia "
    "temporal, y esa sigue siendo débil. Corregida, su lluvia anual se correlaciona con el caudal anual solo entre 0.62 y 0.82, "
    "frente a 0.91–0.94 de CHIRPS, y su variabilidad interanual sigue siendo menor que la observada (CV de 9 % frente a 11 %). El "
    "mapeo de cuantiles vuelve utilizable a ERA5 cuando se necesitan sus otras variables o series en tiempo casi real, pero para el "
    "balance hídrico de esta cuenca no supera a CHIRPS."]:
    P(t, amarillo=True)

d.add_heading("D. Conclusión sobre los productos", 1)
cambio(9, "Conclusión de ERA5 con el resultado del mapeo de cuantiles", "Conclusiones, viñeta de CHIRPS y ERA5: reemplazar las dos "
       "últimas frases.",
       "Un escalamiento lineal mensual podría corregir la media y el ciclo anual. Sin embargo, un quantile mapping con series de "
       "estaciones o un Downscaling podría ser más adecuado para corregir el sesgo.",
       "No deben usarse sin corregir su sesgo. Un mapeo de cuantiles con CHIRPS como referencia corrige su media, su ciclo y su patrón "
       "espacial (cierre de +0.5 % y −2.7 %), pero no su variabilidad interanual: la lluvia corregida se correlaciona con el caudal "
       "anual solo entre 0.62 y 0.82, frente a 0.91–0.94 de CHIRPS. Una corrección más completa requeriría series de estaciones del "
       "IDEAM como referencia o una reducción de escala dinámica (downscaling dinámico) con un modelo regional como WRF "
       "(Posada-Marín et al., 2019).")

d.add_heading("E. Referencia nueva", 1)
P("Agregar en orden alfabético (después de Funk et al., 2015):", cursiva=True, gris=True, tam=10)
P("Gudmundsson, L., Bremnes, J. B., Haugen, J. E. y Engen-Skaugen, T. (2012). Technical Note: Downscaling RCM precipitation to the "
  "station scale using statistical transformations – a comparison of methods. Hydrology and Earth System Sciences, 16, 3383–3390.",
  amarillo=True)

d.save(SALIDA)
print(SALIDA)
