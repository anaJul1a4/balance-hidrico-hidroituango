"""
Punto 4 - Relación temperatura-elevación  T = T0 - α·Z  (Anexo 1).

Criterios de selección (acordados):
  1. Estación dentro de la cuenca.
  2. Tmed, Tmax y Tmin con los 12 meses completos (mismo conjunto para las 3 regresiones).
  3. Coherencia Tmin < Tmed < Tmax.
  4. Ubicación verificable: se excluye Aragón (27015110), cuya coordenada cae a 512 m en el cañón del Cauca
     aunque la estación está a 2600 m en Santa Rosa de Osos (cuenca del Nechí).
  5. Atípicos (|residuo estudentizado| > 2.5) se señalan y se conservan; se reporta la sensibilidad.
Se usa la elevación del catálogo IDEAM (no la del DEM) para Z.

Salidas: trabajo1/resultados/05_temperatura_elevacion/
    Punto4_temperatura_elevacion.xlsx  (fórmulas vivas de Excel)
    diagramas_T_Z.png, diagrama_Tmax/Tmin/Tmed.png, regresiones_mensuales.png
    coeficientes_T_Z.csv (anual) y coeficientes_T_Z_mensual.csv (para los puntos 5 y 6)
"""
import re
from pathlib import Path

import geopandas as gpd
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import rasterio
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
DEL = ROOT / "trabajo1" / "resultados" / "01_delimitacion"
OUT = ROOT / "trabajo1" / "resultados" / "05_temperatura_elevacion"
OUT.mkdir(parents=True, exist_ok=True)
XL = ROOT / "CLIMATOLOGIA_IDEAM" / "Promedios Climatológicos  1981 - 2010.xlsx"
MESES = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", "NOV", "DIC"]
EXCLUIR = {27015110: "Coordenada errónea: cae a 512 m en el cañón del Cauca; la estación (2600 m) está en "
                     "Santa Rosa de Osos, cuenca del Nechí (fuera de la cuenca)"}
UMBRAL_ATIP = 2.5
VARS = {"Tmax": ("Temperatura máxima media", "#d7301f"),
        "Tmin": ("Temperatura mínima media", "#2171b5"),
        "Tmed": ("Temperatura media", "#238b45")}


def dms(s):
    s = str(s).replace(",", ".")
    m = re.findall(r"[\d.]+", s)
    if len(m) < 3:
        return np.nan
    v = float(m[0]) + float(m[1]) / 60 + float(m[2]) / 3600
    return -v if ("W" in s or "S" in s) else v


# ----------------------------------------------------------------------------- inventario y selección
t = pd.read_excel(XL, sheet_name="Temperatura", header=1)
t["lon"], t["lat"] = t.LONGITUD.map(dms), t.LATITUD.map(dms)
t = t.dropna(subset=["lon", "lat"]).copy()
t = t.rename(columns={"ELEVACIÓN": "Z", "ANUAL": "Tmed", "PROMEDIO ANUAL": "Tmax", "PROMEDIO ANUAL.1": "Tmin"})
for m in MESES:
    t[f"Tmed_{m}"], t[f"Tmax_{m}"], t[f"Tmin_{m}"] = t[m], t[m + ".1"], t[m + ".2"]
g = gpd.GeoDataFrame(t, geometry=gpd.points_from_xy(t.lon, t.lat), crs=4326).to_crs("EPSG:3116")
cuenca = gpd.read_file(DEL / "cuenca_cauca_hidroituango.gpkg", layer="cuenca").geometry[0]
g = g[g.within(cuenca)].copy()
with rasterio.open(DEL / "dem_glo30_90m_3116.tif") as s:
    g["Z_DEM"] = [v[0] for v in s.sample(zip(g.geometry.x, g.geometry.y))]
g["dZ"] = g.Z - g.Z_DEM
for v in VARS:
    g[f"meses_{v}"] = g[[f"{v}_{m}" for m in MESES]].notna().sum(1)
g["coherente"] = (g.Tmin < g.Tmed) & (g.Tmed < g.Tmax)


def motivo(r):
    if r.CODIGO in EXCLUIR:
        return EXCLUIR[r.CODIGO]
    falt = [v for v in VARS if r[f"meses_{v}"] < 12]
    if falt:
        return "Serie incompleta (< 12 meses) en " + ", ".join(falt)
    if not r.coherente:
        return "Incoherente: no cumple Tmin < Tmed < Tmax"
    return ""


g["motivo_exclusion"] = g.apply(motivo, axis=1)
g["seleccionada"] = g.motivo_exclusion == ""
inv = g.sort_values("Z")
sel = inv[inv.seleccionada].reset_index(drop=True)
n = len(sel)
print(f"inventario {len(inv)}  seleccionadas {n}")


# ----------------------------------------------------------------------------- cálculos en Python (figuras/verificación)
def regresion(Z, T):
    b, a = np.polyfit(Z, T, 1)
    pred = a + b * Z
    e = T - pred
    sse, sst = np.sum(e ** 2), np.sum((T - T.mean()) ** 2)
    s = np.sqrt(sse / (len(T) - 2))
    h = 1 / len(T) + (Z - Z.mean()) ** 2 / np.sum((Z - Z.mean()) ** 2)
    return dict(a=a, b=b, T0=a, alfa=-b, alfa_km=-b * 1000, R2=1 - sse / sst, n=len(T), s=s,
                pred=pred, e=e, rst=e / (s * np.sqrt(1 - h)))


Z = sel.Z.to_numpy(float)
res = {v: regresion(Z, sel[v].to_numpy(float)) for v in VARS}
sens = []
for v in VARS:
    at = np.abs(res[v]["rst"]) > UMBRAL_ATIP
    sel[f"atipico_{v}"] = at
    r2 = regresion(Z[~at], sel[v].to_numpy(float)[~at]) if at.any() else res[v]
    sens.append({"Variable": v, "Atípicos": ", ".join(sel.NOMBRE[at]) or "ninguno",
                 "n con todos": n, "α con todos (°C/km)": res[v]["alfa_km"], "T0 con todos (°C)": res[v]["T0"],
                 "R² con todos": res[v]["R2"], "n sin atípicos": int((~at).sum()),
                 "α sin atípicos (°C/km)": r2["alfa_km"], "T0 sin atípicos (°C)": r2["T0"],
                 "R² sin atípicos": r2["R2"]})
sens = pd.DataFrame(sens)
coef = pd.DataFrame({v: {k: res[v][k] for k in ["T0", "alfa", "alfa_km", "R2", "n", "s"]} for v in VARS}).T
coef.to_csv(OUT / "coeficientes_T_Z.csv", encoding="utf-8-sig")
mens = []
for v in VARS:
    for m in MESES:
        r = regresion(Z, sel[f"{v}_{m}"].to_numpy(float))
        mens.append({"variable": v, "mes": m, "T0": r["T0"], "alfa": r["alfa"], "alfa_km": r["alfa_km"], "R2": r["R2"]})
mens = pd.DataFrame(mens)
mens.to_csv(OUT / "coeficientes_T_Z_mensual.csv", index=False, encoding="utf-8-sig")
print(coef.round(4).to_string()); print(sens.round(3).to_string())

# ----------------------------------------------------------------------------- diagramas
zz = np.linspace(0, 4500, 50)


def diagrama(ax, v):
    r, (nom, col) = res[v], VARS[v]
    at = sel[f"atipico_{v}"].to_numpy()
    ax.scatter(Z[~at], sel[v][~at], s=28, c=col, edgecolors="black", linewidths=0.4, label="Estaciones IDEAM", zorder=3)
    if at.any():
        ax.scatter(Z[at], sel[v][at], s=60, facecolors="none", edgecolors="black", linewidths=1.2,
                   label=f"Atípico (|r*| > {UMBRAL_ATIP})", zorder=4)
        for zi, ti, nm in zip(Z[at], sel[v][at], sel.NOMBRE[at]):
            ax.annotate(nm, (zi, ti), xytext=(5, 5), textcoords="offset points", fontsize=7)
    ax.plot(zz, r["a"] + r["b"] * zz, c="black", lw=1.3, label="Regresión lineal")
    ax.text(0.97, 0.97,
            f"{v} = {r['T0']:.2f} − {r['alfa']:.5f}·Z\nα = {r['alfa_km']:.2f} °C/km\nR² = {r['R2']:.3f}\nn = {r['n']} estaciones",
            transform=ax.transAxes, ha="right", va="top", fontsize=9,
            bbox=dict(boxstyle="round", fc="white", ec="0.6"))
    ax.set_xlabel("Elevación Z (m s.n.m.)"); ax.set_ylabel(f"{v} (°C)")
    ax.set_title(f"{nom} vs. elevación", fontsize=10)
    ax.set_xlim(0, 4500); ax.set_ylim(0, 36); ax.grid(alpha=0.3); ax.legend(loc="lower left", fontsize=8)


for v in VARS:
    f, a = plt.subplots(figsize=(6.5, 5)); diagrama(a, v); f.tight_layout()
    f.savefig(OUT / f"diagrama_{v}.png", dpi=200); plt.close(f)
f, axs = plt.subplots(1, 3, figsize=(17, 5.2))
for a, v in zip(axs, VARS):
    diagrama(a, v)
f.suptitle("Relación temperatura–elevación, cuenca del río Cauca hasta Hidroituango (IDEAM 1981–2010)",
           fontweight="bold")
f.tight_layout(); f.savefig(OUT / "diagramas_T_Z.png", dpi=200); plt.close(f)

f, axs = plt.subplots(1, 2, figsize=(12, 4))
for v, (nom, col) in VARS.items():
    s = mens[mens.variable == v]
    axs[0].plot(MESES, s.alfa_km, "o-", c=col, label=v); axs[1].plot(MESES, s.R2, "o-", c=col, label=v)
axs[0].set_ylabel("α (°C/km)"); axs[0].set_title("Gradiente térmico vertical mensual")
axs[1].set_ylabel("R²"); axs[1].set_title("Coeficiente de determinación mensual")
for a in axs:
    a.grid(alpha=0.3); a.legend()
f.tight_layout(); f.savefig(OUT / "regresiones_mensuales.png", dpi=200); plt.close(f)

# ----------------------------------------------------------------------------- Excel con fórmulas
wb = Workbook()
H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F4E79")
BOLD = Font(bold=True); thin = Side(style="thin", color="BBBBBB"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
ATF = PatternFill("solid", fgColor="FCE4D6"); EXF = PatternFill("solid", fgColor="F2F2F2")


def encabezado(ws, fila, cols):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(fila, j, c); cell.font = H; cell.fill = HF
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center"); cell.border = BOX


def anchos(ws, w):
    for j, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(j)].width = x


# 1. Criterios
ws = wb.active; ws.title = "Criterios"
txt = [
    ("PUNTO 4 — RELACIÓN TEMPERATURA–ELEVACIÓN", BOLD),
    ("Cuenca del río Cauca hasta la presa de Hidroituango. Fuente: IDEAM, Promedios Climatológicos 1981–2010.", None),
    ("Modelo: T = T0 − α·Z  (regresión T = a + b·Z  ⇒  T0 = a, α = −b; Anexo 1)", None),
    ("", None),
    ("Criterios de selección de estaciones", BOLD),
    ("1. Ubicación dentro de la cuenca delimitada (DEM GLO-30, 90 m).", None),
    ("2. Tmed, Tmax y Tmin con los 12 meses completos; el mismo conjunto de estaciones para las tres regresiones.", None),
    ("3. Coherencia física: Tmin < Tmed < Tmax.", None),
    ("4. Ubicación verificable: se excluye Aragón (27015110); su coordenada cae a 512 m en el cañón del Cauca, "
     "pero la estación está a 2600 m en Santa Rosa de Osos (cuenca del Nechí).", None),
    (f"5. Atípicos: |residuo estudentizado| > {UMBRAL_ATIP}. Se señalan y se conservan; ver hoja 'Sensibilidad_atipicos'.", None),
    ("Z = elevación del catálogo IDEAM (medida en la estación). La hoja 'Inventario' compara con el DEM.", None),
    ("", None),
    (f"Estaciones en la cuenca: {len(inv)}   |   Seleccionadas: {n}   |   Rango de Z: {Z.min():.0f}–{Z.max():.0f} m", BOLD),
    ("", None),
    ("Hojas: Inventario · Datos_anuales (cálculos por estación) · Regresiones (resultados) · Datos_mensuales · "
     "Regresiones_mensuales (insumo ETP Hargreaves) · Sensibilidad_atipicos · Diagramas", None),
    ("Todas las celdas de cálculo son fórmulas de Excel (INTERSECCION.EJE, PENDIENTE, COEFICIENTE.R2, ERROR.TIPICO.XY…).", None),
]
for i, (s, f) in enumerate(txt, 1):
    c = ws.cell(i, 1, s)
    if f:
        c.font = f
ws.column_dimensions["A"].width = 140

# 2. Inventario
ws = wb.create_sheet("Inventario")
cols = ["Código", "Nombre", "Municipio", "Departamento", "Cat.", "Z catálogo (m)", "Z DEM (m)", "Z cat − Z DEM (m)",
        "Meses Tmed", "Meses Tmax", "Meses Tmin", "Tmed (°C)", "Tmax (°C)", "Tmin (°C)", "Seleccionada", "Motivo de exclusión"]
encabezado(ws, 1, cols)
for i, r in enumerate(inv.itertuples(), 2):
    vals = [r.CODIGO, r.NOMBRE, r.MUNICIPIO, r.DEPARTAMENTO, r.CAT, r.Z, round(float(r.Z_DEM), 0), round(float(r.dZ), 0),
            r.meses_Tmed, r.meses_Tmax, r.meses_Tmin, r.Tmed, r.Tmax, r.Tmin, "Sí" if r.seleccionada else "No",
            r.motivo_exclusion]
    for j, x in enumerate(vals, 1):
        c = ws.cell(i, j, None if (isinstance(x, float) and np.isnan(x)) else x); c.border = BOX
        if not r.seleccionada:
            c.fill = EXF
anchos(ws, [11, 20, 18, 16, 6, 10, 10, 11, 7, 7, 7, 9, 9, 9, 11, 90])
ws.freeze_panes = "C2"

# 3. Datos anuales con fórmulas
ws = wb.create_sheet("Datos_anuales")
base = ["Código", "Nombre", "Municipio", "Z (m)", "Tmax (°C)", "Tmin (°C)", "Tmed (°C)", "(Tmax+Tmin)/2 (°C)"]
sub = ["T estimada", "Residuo", "Leverage h", "Residuo estudentizado", "¿Atípico?"]
cols = base + [f"{v} · {s}" for v in VARS for s in sub]
encabezado(ws, 1, cols)
f0, f1 = 2, n + 1
ZR = f"$D${f0}:$D${f1}"
COLT = {"Tmax": "E", "Tmin": "F", "Tmed": "G"}
RROW = {"Tmax": 2, "Tmin": 3, "Tmed": 4}   # fila de cada variable en la hoja Regresiones
for i, r in enumerate(sel.itertuples(), f0):
    for j, x in enumerate([r.CODIGO, r.NOMBRE, r.MUNICIPIO, r.Z, r.Tmax, r.Tmin, r.Tmed], 1):
        ws.cell(i, j, x).border = BOX
    ws.cell(i, 8, f"=(E{i}+F{i})/2").border = BOX
    for k, v in enumerate(VARS):
        c0 = 9 + 5 * k
        L = [get_column_letter(c0 + q) for q in range(5)]
        rr = RROW[v]
        ws.cell(i, c0, f"=Regresiones!$C${rr}+Regresiones!$D${rr}*D{i}")
        ws.cell(i, c0 + 1, f"={COLT[v]}{i}-{L[0]}{i}")
        ws.cell(i, c0 + 2, f"=1/COUNT({ZR})+(D{i}-AVERAGE({ZR}))^2/DEVSQ({ZR})")
        ws.cell(i, c0 + 3, f"={L[1]}{i}/(Regresiones!$J${rr}*SQRT(1-{L[2]}{i}))")
        ws.cell(i, c0 + 4, f'=IF(ABS({L[3]}{i})>{UMBRAL_ATIP},"Sí","No")')
        for q in range(5):
            ws.cell(i, c0 + q).border = BOX
            ws.cell(i, c0 + q).number_format = "0.000"
        if sel.loc[i - f0, f"atipico_{v}"]:
            for q in range(5):
                ws.cell(i, c0 + q).fill = ATF
anchos(ws, [11, 20, 18, 8, 9, 9, 9, 10] + [10] * 15)
ws.freeze_panes = "D2"

# 4. Regresiones (resultados con fórmulas)
ws = wb.create_sheet("Regresiones")
cols = ["Variable", "n", "a = T0 (°C)", "b (°C/m)", "α = −b (°C/m)", "α (°C/km)", "R² (COEFICIENTE.R2)",
        "R² Anexo 1 = 1 − Σe²/Σ(T−T̄)²", "r (Pearson)", "Error típico s (°C)", "Ecuación"]
encabezado(ws, 1, cols)
for v, rr in RROW.items():
    TR = f"Datos_anuales!${COLT[v]}${f0}:${COLT[v]}${f1}"
    ZRx = f"Datos_anuales!{ZR}"
    eL = get_column_letter(9 + 5 * list(VARS).index(v) + 1)
    ER = f"Datos_anuales!${eL}${f0}:${eL}${f1}"
    fila = [v, f"=COUNT({TR})", f"=INTERCEPT({TR},{ZRx})", f"=SLOPE({TR},{ZRx})", f"=-D{rr}", f"=1000*E{rr}",
            f"=RSQ({TR},{ZRx})", f"=1-SUMSQ({ER})/DEVSQ({TR})", f"=CORREL({TR},{ZRx})", f"=STEYX({TR},{ZRx})",
            f'="{v} = "&TEXT(C{rr},"0.00")&" − "&TEXT(E{rr},"0.00000")&"·Z"']
    for j, x in enumerate(fila, 1):
        c = ws.cell(rr, j, x); c.border = BOX
        c.number_format = {3: "0.000", 4: "0.000000", 5: "0.000000", 6: "0.00", 7: "0.000", 8: "0.000",
                           9: "0.000", 10: "0.000"}.get(j, "General")
ws.cell(6, 1, "Verificación (calculado en Python con numpy.polyfit)").font = BOLD
encabezado(ws, 7, ["Variable", "n", "T0 (°C)", "α (°C/km)", "R²"])
for i, v in enumerate(VARS, 8):
    for j, x in enumerate([v, res[v]["n"], round(res[v]["T0"], 3), round(res[v]["alfa_km"], 3), round(res[v]["R2"], 4)], 1):
        ws.cell(i, j, x).border = BOX
anchos(ws, [10, 6, 12, 12, 13, 10, 13, 18, 11, 13, 32])

# 5. Datos mensuales
ws = wb.create_sheet("Datos_mensuales")
cols = ["Código", "Nombre", "Z (m)"] + [f"{v} {m}" for v in VARS for m in MESES]
encabezado(ws, 1, cols)
for i, r in enumerate(sel.itertuples(), f0):
    vals = [r.CODIGO, r.NOMBRE, r.Z] + [getattr(r, f"{v}_{m}") for v in VARS for m in MESES]
    for j, x in enumerate(vals, 1):
        ws.cell(i, j, x).border = BOX
anchos(ws, [11, 20, 8] + [8] * 36)
ws.freeze_panes = "D2"

# 6. Regresiones mensuales
ws = wb.create_sheet("Regresiones_mensuales")
encabezado(ws, 1, ["Variable", "Mes", "n", "T0 (°C)", "α (°C/m)", "α (°C/km)", "R²"])
fila = 2
for k, v in enumerate(VARS):
    for mi, m in enumerate(MESES):
        cl = get_column_letter(4 + 12 * k + mi)
        TR = f"Datos_mensuales!${cl}${f0}:${cl}${f1}"
        ZRm = f"Datos_mensuales!$C${f0}:$C${f1}"
        for j, x in enumerate([v, m, f"=COUNT({TR})", f"=INTERCEPT({TR},{ZRm})", f"=-SLOPE({TR},{ZRm})",
                               f"=1000*E{fila}", f"=RSQ({TR},{ZRm})"], 1):
            c = ws.cell(fila, j, x); c.border = BOX
            c.number_format = {4: "0.000", 5: "0.000000", 6: "0.00", 7: "0.000"}.get(j, "General")
        fila += 1
ws.add_image(XLImage(str(OUT / "regresiones_mensuales.png")), "I2")
anchos(ws, [9, 6, 5, 10, 12, 10, 8])

# 7. Sensibilidad
ws = wb.create_sheet("Sensibilidad_atipicos")
encabezado(ws, 1, list(sens.columns))
for i, r in enumerate(sens.itertuples(index=False), 2):
    for j, x in enumerate(r, 1):
        c = ws.cell(i, j, round(x, 4) if isinstance(x, float) else x); c.border = BOX
anchos(ws, [9, 40, 10, 14, 14, 11, 12, 14, 14, 12])

# 8. Diagramas
ws = wb.create_sheet("Diagramas")
for k, v in enumerate(VARS):
    img = XLImage(str(OUT / f"diagrama_{v}.png")); img.width, img.height = 585, 450
    ws.add_image(img, f"A{1 + 24 * k}")

wb.save(OUT / "Punto4_temperatura_elevacion.xlsx")
print("Excel guardado")
