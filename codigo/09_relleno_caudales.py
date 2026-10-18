"""
Punto 8 (complemento) - Revisión de faltantes y relleno de caudales mensuales 1981-2010.

Revisión: en los archivos IDEAM (Q_MEDIA_D) el faltante es la ausencia de la fecha; se verifican ceros y negativos.
Relleno: regresión entre estaciones con caudales mensuales en logaritmo
        ln Q_obj = a + b · ln Q_pred      (corrección de sesgo de Duan al retransformar)
  · Predictoras: las 102 estaciones con Q_MEDIA_D dentro de la cuenca (cauce principal, tributarios y cabecera).
  · Una predictora es elegible si comparte ≥ 60 meses con la estación objetivo en 1981-2010 y r(ln) ≥ 0.80.
  · Cascada: cada mes faltante se rellena con la predictora de mayor r que tenga dato ese mes.
  · Validación cruzada dejando un año fuera: se reconstruyen los meses observados con la misma cascada,
    ajustando las regresiones sin el año evaluado.
Estaciones objetivo: las 10 validadas + Pte. Pescadero y Olaya (no validadas por < 15 años completos).

Requiere 08_excel_caudales.py. Escribe el relleno en un libro aparte (Punto8_caudales_relleno.xlsx) y sólo
agrega la hoja 'Faltantes' a Punto8_caudales_estaciones.xlsx.
"""
import json
import sys
from pathlib import Path

import numpy as np
np.in1d = np.isin

import geopandas as gpd
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pyproj import Transformer
from pysheds.grid import Grid
from pysheds.view import Raster

sys.path.insert(0, str(Path(__file__).parent))
from priority_flood import condicionar  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
DAT = ROOT / "trabajo1" / "datos" / "caudales_ideam"
DEL = ROOT / "trabajo1" / "resultados" / "01_delimitacion"
OUT = ROOT / "trabajo1" / "resultados" / "06_caudales"
XLSX = OUT / "Punto8_caudales_estaciones.xlsx"
P0, P1 = "1981", "2010"
MIN_DIAS, N_MIN, R_MIN = 0.8, 60, 0.80
EXTRA = {"26237050": "PUENTE PESCADERO", "26237100": "OLAYA"}
SEG_ANO = 365.25 * 86400
DIRMAP = (64, 128, 1, 2, 4, 8, 16, 32)

cat = pd.DataFrame(json.load(open(ROOT / "trabajo1" / "datos" / "catalogo_ideam_zona26_hidrometricas.json",
                                  encoding="utf-8"))).drop_duplicates("codigo")
cat["cod"] = cat.codigo.str[2:]
cat = cat.set_index("cod")
nombre = lambda c: cat.nombre.get(c, c).split(" [")[0].replace(" - AUT", "").strip().title()

# ----------------------------------------------------------------------------- lectura y revisión de faltantes
rev, M = [], {}
dias_per = pd.date_range(f"{P0}-01-01", f"{P1}-12-31")
for f in sorted((DAT / "q_media_diaria").glob("Q_MEDIA_D@*.data")):
    c = f.stem.split("@")[1]
    d = pd.read_csv(f, sep="|", parse_dates=["Fecha"]).set_index("Fecha").Valor.astype(float)
    d = d[~d.index.duplicated()]
    p = d[P0:P1]
    rev.append({"codigo": c, "ceros": int((p == 0).sum()), "negativos": int((p < 0).sum()),
                "dias_con_dato": int(p.notna().sum()), "dias_faltantes": len(dias_per) - int(p.reindex(dias_per).notna().sum())})
    d[d <= 0] = np.nan                              # ceros/negativos se tratarían como faltantes (no hay ninguno)
    cnt = d.resample("MS").count()
    M[c] = d.resample("MS").mean().where(cnt >= MIN_DIAS * cnt.index.days_in_month)
M = pd.DataFrame(M).loc[P0:P1].reindex(pd.date_range(f"{P0}-01-01", f"{P1}-12-01", freq="MS"))
rev = pd.DataFrame(rev).set_index("codigo")
L = np.log(M)

disp = pd.read_csv(DAT / "disponibilidad_caudales.csv", dtype={"codigo": str})
objetivos = list(disp.loc[disp["apta (>=15 anos 1981-2010)"] == True, "codigo"]) + list(EXTRA)


# ----------------------------------------------------------------------------- regresiones
def ajustar(y, x):
    ok = y.notna() & x.notna()
    if ok.sum() < N_MIN:
        return None
    b, a = np.polyfit(x[ok], y[ok], 1)
    e = y[ok] - (a + b * x[ok])
    return dict(a=a, b=b, r=np.corrcoef(x[ok], y[ok])[0, 1], n=int(ok.sum()), duan=float(np.mean(np.exp(e))))


def cascada(t, excluir_ano=None):
    """Predictoras elegibles ordenadas por r; si excluir_ano, se ajustan sin ese año."""
    y = L[t].copy()
    if excluir_ano is not None:
        y[y.index.year == excluir_ano] = np.nan
    regs = []
    for p in L.columns:
        if p == t:
            continue
        rg = ajustar(y, L[p])
        if rg and rg["r"] >= R_MIN:
            regs.append((p, rg))
    return sorted(regs, key=lambda z: -z[1]["r"])


relleno, fuente, regs_usadas, cv = {}, {}, [], []
for t in objetivos:
    regs = cascada(t)
    q, src = M[t].copy(), pd.Series("observado", index=M.index, dtype=object)
    for p, rg in regs:
        hueco = q.isna() & M[p].notna()
        if hueco.any():
            q[hueco] = np.exp(rg["a"] + rg["b"] * L[p][hueco]) * rg["duan"]
            src[hueco] = p
            regs_usadas.append({"codigo_obj": t, "estacion_obj": nombre(t), "predictora": p, "nombre_pred": nombre(p),
                                "corriente_pred": cat.corriente.get(p, ""), "n_traslape": rg["n"], "r_ln": rg["r"],
                                "a": rg["a"], "b": rg["b"], "factor_Duan": rg["duan"], "meses_rellenados": int(hueco.sum())})
    relleno[t], fuente[t] = q, src
    # validación cruzada dejando un año fuera
    obs, est = [], []
    for ano in sorted(set(M.index.year[M[t].notna()])):
        sel = (M.index.year == ano) & M[t].notna()
        pred = pd.Series(np.nan, index=M.index[sel])
        for p, rg in cascada(t, excluir_ano=ano):
            h = pred.isna() & M[p][sel].notna().to_numpy()
            pred[h] = np.exp(rg["a"] + rg["b"] * L[p][sel][h.to_numpy()]) * rg["duan"]
        ok = pred.notna()
        obs += list(M[t][sel][ok.to_numpy()]); est += list(pred[ok])
    obs, est = np.array(obs), np.array(est)
    cv.append({"codigo": t, "n_validacion": len(obs),
               "NSE": 1 - np.sum((est - obs) ** 2) / np.sum((obs - obs.mean()) ** 2),
               "RMSE (m3/s)": np.sqrt(np.mean((est - obs) ** 2)), "Sesgo (%)": 100 * (est.mean() - obs.mean()) / obs.mean(),
               "r": np.corrcoef(obs, est)[0, 1]})
R = pd.DataFrame(relleno)
F = pd.DataFrame(fuente)
regs_usadas = pd.DataFrame(regs_usadas)
cv = pd.DataFrame(cv).set_index("codigo")

# ----------------------------------------------------------------------------- áreas de Pescadero y Olaya (DEM)
areas = gpd.read_file(OUT / "subcuencas_estaciones.gpkg", layer="estaciones").set_index("codigo").area_km2_DEM.to_dict()
grid = Grid.from_raster(str(DEL / "dem_glo30_90m_3116.tif")); dem = grid.read_raster(str(DEL / "dem_glo30_90m_3116.tif"))
fdir = grid.flowdir(Raster(condicionar(dem), viewfinder=dem.viewfinder), dirmap=DIRMAP)
acc = grid.accumulation(fdir, dirmap=DIRMAP); ck = abs(grid.affine.a * grid.affine.e) / 1e6
tr = Transformer.from_crs(4326, "EPSG:3116", always_xy=True)
for c in EXTRA:
    x, y = tr.transform(float(cat.longitud[c]), float(cat.latitud[c]))
    xs, ys = grid.snap_to_mask(acc > 1000 / ck, (x, y))
    areas[c] = float(np.asarray(grid.catchment(x=xs, y=ys, fdir=fdir, dirmap=DIRMAP, xytype="coordinate")).sum() * ck)
A_CUENCA = float(gpd.read_file(DEL / "cuenca_cauca_hidroituango.gpkg", layer="cuenca").area_km2[0])

# ----------------------------------------------------------------------------- comparación de medias
comp = []
for t in objetivos:
    o = M[t]
    anual_o = o.groupby(o.index.year).apply(lambda s: s.mean() if s.count() == 12 else np.nan)
    anual_r = R[t].groupby(R.index.year).apply(lambda s: s.mean() if s.count() == 12 else np.nan)
    comp.append({"codigo": t, "estacion": nombre(t), "area_km2": areas[t],
                 "meses_observados": int(o.notna().sum()), "meses_rellenados": int((F[t] != "observado").sum() - R[t].isna().sum()),
                 "meses_sin_rellenar": int(R[t].isna().sum()),
                 "Q años completos observados": anual_o.mean(), "años completos observados": int(anual_o.count()),
                 "Q meses válidos observados": o.mean(),
                 "Q serie rellenada": anual_r.mean(), "años serie rellenada": int(anual_r.count())})
comp = pd.DataFrame(comp).set_index("codigo").join(cv)
comp["dif. rellenada vs. obs. (%)"] = 100 * (comp["Q serie rellenada"] / comp["Q años completos observados"] - 1)
comp["R serie rellenada (mm/año)"] = comp["Q serie rellenada"] * SEG_ANO / (comp.area_km2 * 1e6) * 1000
pd.set_option("display.width", 250)
print(comp.round(2).to_string())
print(regs_usadas.groupby("codigo_obj").apply(lambda d: ", ".join(f"{r.nombre_pred}({r.meses_rellenados})" for r in d.itertuples())).to_string())
for c in ["26237050", "26237040"]:
    print(f"Hidroituango desde {nombre(c)}: {comp.loc[c, 'Q serie rellenada'] * A_CUENCA / comp.loc[c, 'area_km2']:.1f} m3/s")

# ----------------------------------------------------------------------------- figura
fig, axs = plt.subplots(len(objetivos), 1, figsize=(13, 1.35 * len(objetivos)), sharex=True)
for ax, t in zip(axs, objetivos):
    ax.plot(R.index, R[t], c="#d73027", lw=0.9, label="rellenado")
    ax.plot(M.index, M[t], c="#08306b", lw=0.9, label="observado")
    ax.set_ylabel(nombre(t), rotation=0, ha="right", va="center", fontsize=8); ax.tick_params(labelsize=7)
    ax.text(0.995, 0.92, f"NSE val. = {cv.loc[t, 'NSE']:.2f}", transform=ax.transAxes, ha="right", va="top", fontsize=7)
axs[0].legend(loc="upper left", fontsize=7, ncol=2)
axs[0].set_title("Caudal medio mensual (m³/s) 1981–2010: observado y rellenado por regresión entre estaciones", fontsize=10)
fig.tight_layout(); fig.savefig(OUT / "relleno_caudales.png", dpi=180); plt.close(fig)

# ----------------------------------------------------------------------------- Excel
wb = load_workbook(XLSX)
for s in ["Faltantes", "Relleno_metodo", "Relleno_regresiones", "Q_mensual_relleno", "Q_fuente_relleno",
          "Q_anual_relleno", "Comparacion_Q"]:
    if s in wb.sheetnames:
        del wb[s]
principal = wb                 # Excel del punto 8 (datos observados): sólo recibe la hoja 'Faltantes'
wb = Workbook()                # el relleno va en un Excel aparte, por si se descarta
wb.remove(wb.active)
XLSX_RELLENO = OUT / "Punto8_caudales_relleno.xlsx"
H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F4E79"); BOLD = Font(bold=True)
thin = Side(style="thin", color="BBBBBB"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
RF = PatternFill("solid", fgColor="F8CBAD")


def encabezado(ws, cols, alto=45):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(1, j, c); cell.font = H; cell.fill = HF; cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.row_dimensions[1].height = alto


def tabla(ws, df, anchos=None, fmt="0.00"):
    encabezado(ws, list(df.columns))
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, x in enumerate(row, 1):
            c = ws.cell(i, j, None if (isinstance(x, float) and np.isnan(x)) else x); c.border = BOX
            if isinstance(x, float):
                c.number_format = fmt
    for j, w in enumerate(anchos or [14] * len(df.columns), 1):
        ws.column_dimensions[get_column_letter(j)].width = w


# Faltantes (en ambos libros)
fal = rev.loc[objetivos].assign(estacion=[nombre(c) for c in objetivos])
fal["% días faltantes"] = 100 * fal.dias_faltantes / len(dias_per)
fal["meses válidos (de 360)"] = [int(M[c].notna().sum()) for c in objetivos]
fal = fal.reset_index()[["codigo", "estacion", "dias_con_dato", "dias_faltantes", "% días faltantes", "ceros",
                         "negativos", "meses válidos (de 360)"]]
for libro, filas in [(principal, fal[~fal.codigo.isin(EXTRA)]), (wb, fal)]:
    ws = libro.create_sheet("Faltantes")
    tabla(ws, filas, [10, 18, 12, 12, 11, 8, 10, 12], "0.0")
    ws.cell(len(filas) + 3, 1, "En los archivos IDEAM el dato faltante es la ausencia de la fecha (no se codifica con 0 "
            "ni −999). No hay ceros ni negativos en 1981–2010; si los hubiera se tratarían como faltantes.")
    ws.cell(len(filas) + 4, 1, "Tratamiento: los faltantes NO se rellenan; un mes es válido con ≥ 80 % de días y el caudal "
            "medio multianual usa sólo años con 12 meses válidos." if libro is principal else
            "En este libro los meses faltantes se rellenan por regresión entre estaciones (ver Relleno_metodo).")

# Método
ws = wb.create_sheet("Relleno_metodo")
for i, s in enumerate([
        "RELLENO DE CAUDALES MENSUALES 1981–2010 — REGRESIÓN ENTRE ESTACIONES",
        "Modelo: ln Q_objetivo = a + b · ln Q_predictora ;  Q = exp(a + b·ln Q_pred) · factor de Duan (corrección de sesgo).",
        f"Predictoras posibles: 102 estaciones IDEAM con caudal medio diario dentro de la cuenca (cauce principal, tributarios y cabecera).",
        f"Elegibilidad: ≥ {N_MIN} meses de traslape en 1981–2010 y r(ln Q) ≥ {R_MIN}.",
        "Cascada: cada mes faltante se rellena con la predictora elegible de mayor r que tenga dato ese mes.",
        "Validación cruzada: dejando un año fuera, se reconstruyen los meses observados con la misma cascada (NSE, RMSE, sesgo en 'Comparacion_Q').",
        "Mes válido = ≥ 80 % de los días con dato. Celdas rosadas en 'Q_mensual_relleno' = meses rellenados; la hoja 'Q_fuente_relleno' indica la predictora usada.",
        "Se incluyen Pte. Pescadero y Olaya (no validadas por < 15 años completos) para aprovechar su cercanía al punto de cierre.",
        "El relleno no usa P − ETR, para no condicionar la validación del balance hídrico (punto 9)."], 1):
    ws.cell(i, 1, s).font = BOLD if i == 1 else Font()
ws.column_dimensions["A"].width = 150

ws = wb.create_sheet("Relleno_regresiones")
tabla(ws, regs_usadas, [10, 16, 10, 22, 16, 9, 8, 9, 9, 9, 10], "0.000")

cols_t = [f"{nombre(t)} ({t})" for t in objetivos]
ws = wb.create_sheet("Q_mensual_relleno"); encabezado(ws, ["Fecha"] + cols_t, 40)
ws2 = wb.create_sheet("Q_fuente_relleno"); encabezado(ws2, ["Fecha"] + cols_t, 40)
for i, t in enumerate(R.index, 2):
    ws.cell(i, 1, t.to_pydatetime()).number_format = "yyyy-mm"
    ws2.cell(i, 1, t.to_pydatetime()).number_format = "yyyy-mm"
    for j, c in enumerate(objetivos, 2):
        v = R.at[t, c]
        cell = ws.cell(i, j, None if pd.isna(v) else round(float(v), 2)); cell.number_format = "0.0"
        s = F.at[t, c]
        ws2.cell(i, j, "observado" if s == "observado" else ("sin dato" if pd.isna(v) else f"{nombre(s)} ({s})"))
        if s != "observado" and not pd.isna(v):
            cell.fill = RF
for w in (ws, ws2):
    w.column_dimensions["A"].width = 9
    for j in range(2, len(objetivos) + 2):
        w.column_dimensions[get_column_letter(j)].width = 15
    w.freeze_panes = "B2"

ws = wb.create_sheet("Q_anual_relleno"); encabezado(ws, ["Año"] + cols_t, 40)
for k, ano in enumerate(range(int(P0), int(P1) + 1)):
    i = k + 2; ws.cell(i, 1, ano)
    for j in range(len(objetivos)):
        Lc = get_column_letter(2 + j); rng = f"Q_mensual_relleno!{Lc}{2 + 12 * k}:{Lc}{13 + 12 * k}"
        ws.cell(i, 2 + j, f'=IF(COUNT({rng})=12,AVERAGE({rng}),"")').number_format = "0.0"
nf = int(P1) - int(P0) + 2

ws = wb.create_sheet("Comparacion_Q")
cols = ["Código", "Estación", "Área (km²)", "Meses observados", "Meses rellenados", "Meses sin rellenar",
        "Q años completos observados (m³/s)", "Años completos observados", "Q meses válidos observados (m³/s)",
        "Q serie rellenada 1981–2010 (m³/s)", "Años serie rellenada", "Diferencia rellenada vs. observada (%)",
        "Lámina R serie rellenada (mm/año)", "NSE validación", "RMSE validación (m³/s)", "Sesgo validación (%)"]
encabezado(ws, cols, 60)
for j, t in enumerate(objetivos):
    i = j + 2; r = comp.loc[t]; Lc = get_column_letter(2 + j)
    vals = [t, r.estacion, round(r.area_km2, 1), r.meses_observados, r.meses_rellenados, r.meses_sin_rellenar,
            round(r["Q años completos observados"], 2), r["años completos observados"], round(r["Q meses válidos observados"], 2),
            f"=AVERAGE(Q_anual_relleno!{Lc}2:{Lc}{nf})", f"=COUNT(Q_anual_relleno!{Lc}2:{Lc}{nf})",
            f"=(J{i}/G{i}-1)*100", f"=J{i}*{SEG_ANO:.0f}/(C{i}*1000000)*1000",
            round(r.NSE, 3), round(r["RMSE (m3/s)"], 1), round(r["Sesgo (%)"], 2)]
    for k, x in enumerate(vals, 1):
        c = ws.cell(i, k, x); c.border = BOX
        if k in (10, 12, 13):
            c.number_format = "0.0"
fila = len(objetivos) + 3
ws.cell(fila, 1, "Caudal medio en la presa de Hidroituango (transferencia por relación de áreas, serie rellenada)").font = BOLD
for k, c in enumerate(["26237050", "26237040"]):
    i = objetivos.index(c) + 2
    ws.cell(fila + 1 + k, 1, f"Desde {nombre(c)}").border = BOX
    cell = ws.cell(fila + 1 + k, 2, f"=J{i}*{A_CUENCA:.1f}/C{i}"); cell.border = BOX; cell.number_format = "0.0"
for j, w in enumerate([10, 17, 10, 9, 9, 9, 12, 10, 12, 12, 9, 12, 11, 9, 10, 9], 1):
    ws.column_dimensions[get_column_letter(j)].width = w
from openpyxl.drawing.image import Image as XLImage  # noqa: E402
img = XLImage(str(OUT / "relleno_caudales.png")); img.width, img.height = 1000, 1000 * 1.35 * len(objetivos) / 13
ws.add_image(img, f"A{fila + 5}")
wb.save(XLSX_RELLENO)
principal.save(XLSX)
R.to_csv(DAT / "Q_mensual_1981_2010_rellenado.csv", encoding="utf-8-sig")
print("Excel actualizado")
